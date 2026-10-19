import os
import io
import zipfile
from typing import Tuple, List, Dict, Any
import pdfplumber
import docx
from backend.core.schemas import DocumentMetadata

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50 MB

class IngestionError(Exception):
    """Custom exception raised when file validation or reading fails."""
    pass

def validate_file_signature(content: bytes, filename: str) -> str:
    """
    Validates magic bytes / real file signatures:
    - PDF must start with '%PDF-'
    - DOCX must be a ZIP container containing 'word/document.xml'
    """
    if len(content) == 0:
        raise IngestionError("File is empty (0 bytes).")
    
    if len(content) > MAX_FILE_SIZE:
        raise IngestionError(f"File exceeds maximum allowed size of 50 MB ({len(content)} bytes).")

    # Check PDF magic bytes (%PDF-)
    if content.startswith(b"%PDF-"):
        return "pdf"

    # Check DOCX (ZIP format containing word/document.xml)
    if content.startswith(b"PK\x03\x04") or content.startswith(b"PK\x05\x06") or content.startswith(b"PK\x07\x08"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                namelist = zf.namelist()
                if "word/document.xml" in namelist:
                    return "docx"
        except zipfile.BadZipFile:
            raise IngestionError("Corrupted DOCX/ZIP file archive.")

    raise IngestionError(
        f"Unsupported or invalid file format for '{filename}'. PaperLens only accepts valid PDF (%PDF-) and DOCX files."
    )

def extract_text_from_pdf(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Extracts text page by page using pdfplumber, handling multi-column text layouts
    and tracking page metadata.
    """
    pages_data = []
    full_text_parts = []
    
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if not pdf.pages:
                raise IngestionError("PDF contains no readable pages.")
            
            for page_idx, page in enumerate(pdf.pages, start=1):
                # pdfplumber layout extraction
                page_text = page.extract_text(layout=True) or page.extract_text() or ""
                # Clean up null bytes or irregular controls
                clean_page_text = page_text.replace("\x00", "")
                pages_data.append({
                    "page_number": page_idx,
                    "text": clean_page_text,
                    "char_count": len(clean_page_text)
                })
                if clean_page_text.strip():
                    full_text_parts.append(clean_page_text)
    except Exception as e:
        if isinstance(e, IngestionError):
            raise e
        raise IngestionError(f"Failed to parse PDF document: {str(e)}")

    full_text = "\n\n".join(full_text_parts)
    if not full_text.strip():
        raise IngestionError("Extracted text is empty. The PDF may be scanned/image-only.")

    return full_text, pages_data

def extract_text_from_docx(content: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Extracts text from DOCX paragraphs and tables.
    """
    try:
        doc = docx.Document(io.BytesIO(content))
        paragraphs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    paragraphs.append(" | ".join(row_cells))
        
        full_text = "\n\n".join(paragraphs)
        if not full_text.strip():
            raise IngestionError("DOCX document contains no text content.")
        
        pages_data = [{
            "page_number": 1,
            "text": full_text,
            "char_count": len(full_text)
        }]
        return full_text, pages_data
    except Exception as e:
        if isinstance(e, IngestionError):
            raise e
        raise IngestionError(f"Failed to parse DOCX document: {str(e)}")

def ingest_document(content: bytes, filename: str) -> Tuple[str, List[Dict[str, Any]], DocumentMetadata]:
    """
    End-to-end ingestion orchestrator: validates signature, extracts text,
    and returns full text, per-page data, and metadata.
    """
    file_type = validate_file_signature(content, filename)
    
    if file_type == "pdf":
        full_text, pages_data = extract_text_from_pdf(content)
    elif file_type == "docx":
        full_text, pages_data = extract_text_from_docx(content)
    else:
        raise IngestionError(f"Unsupported file type: {file_type}")

    metadata = DocumentMetadata(
        filename=filename,
        file_type=file_type,
        file_size_bytes=len(content),
        total_pages=len(pages_data),
        total_characters=len(full_text),
        total_words=len(full_text.split())
    )

    return full_text, pages_data, metadata
